import base64
import math
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from citation_utils import format_reference, reference_sources, replace_citation_tokens

BASE_DIR = Path(__file__).resolve().parent
LOGO_PATH = BASE_DIR / "assets" / "logo_ums.png"

def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.3)

def font_run(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_page_number(section, fmt, start=1):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))

def page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, end])

def toc_field(paragraph):
    r = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, sep])
    rr = paragraph.add_run("Daftar isi akan diperbarui saat dokumen dibuka di Microsoft Word. Jika nomor halaman belum sesuai, tekan Ctrl+A lalu F9.")
    font_run(rr, 10, italic=True)
    paragraph.add_run()._r.append(end)

def update_fields(doc):
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")

def clear_footer(section):
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        p.clear()

def remove_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")

def set_cell_fill(cell, fill="E8E7FF"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def add_center(doc, text, bold=False, after=0, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    font_run(r, size, bold=bold)
    return p

def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")

def add_paragraphs(doc, text):
    """Every blank-line-separated block becomes its own academic paragraph."""
    for block in [x.strip() for x in (text or "").split("\n\n") if x.strip()]:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Cm(1.25)
        r = p.add_run(block)
        font_run(r)

def add_numbered(doc, items):
    for i, item in enumerate(items or [], 1):
        if not str(item).strip():
            continue
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(.75)
        p.paragraph_format.first_line_indent = Cm(-.75)
        r = p.add_run(f"{i}. {str(item).strip()}")
        font_run(r)

def add_content_list(doc, items, ordered=True):
    style = "List Number" if ordered else "List Bullet"
    for item in items or []:
        if not str(item).strip():
            continue
        paragraph = doc.add_paragraph(style=style)
        font_run(paragraph.add_run(str(item).strip()))

def add_content_table(doc, text):
    rows = [[cell.strip() for cell in line.split("|")] for line in (text or "").splitlines() if line.strip()]
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    headers = (rows[0] + [""] * column_count)[:column_count]
    body = [(row + [""] * column_count)[:column_count] for row in rows[1:]]
    add_simple_table(doc, headers, body)


def add_content_chart(doc, block):
    lines = [line.strip() for line in (block.get("chartData") or "").splitlines() if line.strip()]
    chart_type = block.get("chartType", "bar")
    comparison = chart_type in ("bar-comparison", "line-comparison")
    if comparison:
        if len(lines) < 2:
            raise ValueError("Grafik perbandingan memerlukan header dan minimal satu baris data.")
        headers = [cell.strip() for cell in lines[0].split("|")]
        if len(headers) < 3 or not headers[0] or any(not label for label in headers[1:]):
            raise ValueError("Header grafik perbandingan harus berformat Kategori | Seri 1 | Seri 2 (atau lebih).")
        if len(headers) > 7:
            raise ValueError("Grafik perbandingan maksimal memiliki 6 seri.")
        series_names = headers[1:]
        rows = []
        for line in lines[1:]:
            parts = [cell.strip() for cell in line.split("|")]
            if len(parts) != len(headers) or not parts[0]:
                raise ValueError("Setiap baris harus memiliki kategori dan nilai untuk semua seri.")
            try:
                values = [float(value.replace(",", ".")) for value in parts[1:]]
            except ValueError as error:
                raise ValueError(f"Semua nilai grafik untuk kategori '{parts[0]}' harus berupa angka.") from error
            if any(not math.isfinite(value) or value < 0 for value in values):
                raise ValueError("Nilai grafik harus angka positif atau nol.")
            rows.append((parts[0], values))
    else:
        series_names = ["Nilai"]
        rows = []
        for line in lines:
            parts = [cell.strip() for cell in line.split("|")]
            if len(parts) != 2 or not parts[0]:
                raise ValueError("Data grafik harus berformat Kategori | Nilai, satu baris per data.")
            try:
                value = float(parts[1].replace(",", "."))
            except ValueError as error:
                raise ValueError(f"Nilai grafik untuk '{parts[0]}' harus berupa angka.") from error
            if not math.isfinite(value) or value < 0:
                raise ValueError("Nilai grafik harus angka positif atau nol.")
            rows.append((parts[0], [value]))

    if not rows:
        raise ValueError("Grafik belum memiliki data. Isi minimal satu baris kategori dan nilai.")
    if len(rows) > 30:
        raise ValueError("Grafik maksimal memuat 30 kategori.")
    if chart_type == "pie" and any(values[0] <= 0 for _, values in rows):
        raise ValueError("Semua nilai pada grafik lingkaran harus lebih besar dari nol.")

    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    labels = [label for label, _ in rows]
    series_values = [[values[index] for _, values in rows] for index in range(len(series_names))]
    colors = ["#635bff", "#ec4899", "#18a976", "#f59e0b", "#38bdf8", "#a855f7"]
    figure = Figure(figsize=(8, 4.5), dpi=160, facecolor="white")
    axis = figure.add_subplot(1, 1, 1)
    title = (block.get("chartTitle") or "").strip()
    if title:
        axis.set_title(title, fontsize=13, pad=14)

    if chart_type == "pie":
        values = [row_values[0] for row_values in series_values]
        axis.pie(values, labels=labels, autopct="%1.1f%%", startangle=90)
        axis.axis("equal")
    elif chart_type in ("line", "line-comparison"):
        for index, (series_name, values) in enumerate(zip(series_names, series_values)):
            axis.plot(
                labels,
                values,
                marker="o",
                linewidth=2.2,
                color=colors[index % len(colors)],
                label=series_name if comparison else None,
            )
        axis.set_ylabel("Nilai")
        axis.grid(axis="y", alpha=0.25)
        if comparison:
            axis.legend()
    elif chart_type == "bar-comparison":
        import numpy as np

        positions = np.arange(len(labels))
        width = min(0.8 / len(series_names), 0.32)
        for index, (series_name, values) in enumerate(zip(series_names, series_values)):
            offset = (index - (len(series_names) - 1) / 2) * width
            axis.bar(
                positions + offset,
                values,
                width,
                color=colors[index % len(colors)],
                label=series_name,
            )
        axis.set_xticks(positions, labels)
        axis.legend()
        axis.set_ylabel("Nilai")
        axis.grid(axis="y", alpha=0.25)
        axis.set_axisbelow(True)
    else:
        bars = axis.bar(labels, series_values[0], color=colors[0], width=0.65)
        axis.bar_label(bars, padding=3, fontsize=8)
        axis.set_ylabel("Nilai")
        axis.grid(axis="y", alpha=0.25)
        axis.set_axisbelow(True)

    axis.tick_params(axis="x", labelrotation=20, labelsize=8)
    figure.tight_layout()
    image_stream = BytesIO()
    FigureCanvasAgg(figure).print_png(image_stream)
    image_stream.seek(0)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(image_stream, width=Cm(13))
    caption = (block.get("caption") or title).strip()
    if caption:
        caption_paragraph = doc.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_run(caption_paragraph.add_run(caption), 10, italic=True)


def add_content_block(doc, block, sources=None, citation_style="apa", citation_numbers=None):
    block_type = block.get("type")
    if block_type == "paragraph":
        text = replace_citation_tokens(
            block.get("text", ""),
            sources or [],
            citation_style,
            citation_numbers or {},
        )
        add_paragraphs(doc, text)
    elif block_type == "list":
        add_content_list(doc, block.get("items", []), block.get("ordered", True))
    elif block_type == "table":
        add_content_table(doc, block.get("tableText", ""))
    elif block_type == "chart":
        add_content_chart(doc, block)
    elif block_type == "image":
        data_url = block.get("dataUrl", "")
        if not data_url.startswith("data:image/") or "," not in data_url:
            return
        try:
            image_data = base64.b64decode(data_url.split(",", 1)[1], validate=True)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture = paragraph.add_run().add_picture(BytesIO(image_data), width=Cm(13))
            max_height = Cm(21)
            if picture.height > max_height:
                scale = max_height / picture.height
                picture.width = int(picture.width * scale)
                picture.height = int(picture.height * scale)
        except Exception:
            return
        caption = (block.get("caption") or "").strip()
        if caption:
            caption_paragraph = doc.add_paragraph()
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font_run(caption_paragraph.add_run(caption), 10, italic=True)

def add_simple_table(doc, headers, rows, widths_font=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c = table.cell(0,i)
        c.text = str(h)
        set_cell_fill(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                font_run(r, widths_font, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text = str(val or "")
            for p in cells[i].paragraphs:
                p.paragraph_format.line_spacing = 1
                for r in p.runs:
                    font_run(r, widths_font)
    return table

def build_docx(data):
    ident = data.get("identitas", {})
    b1 = data.get("bab1", {})
    b2 = data.get("bab2", {})
    b3 = data.get("bab3", {})
    refs = data.get("referensi", [])
    sources = data.get("sources", [])
    citation_style = "ieee" if data.get("citationStyle") == "ieee" else "apa"
    bibliography_sources = reference_sources(data.get("chapters", []), sources)
    citation_numbers = {
        source.get("id"): index
        for index, source in enumerate(bibliography_sources, 1)
        if source.get("id")
    }

    doc = Document()
    configure_section(doc.sections[0])

    # Styles
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for sname in ("Heading 1","Heading 2","Heading 3"):
        st = doc.styles[sname]
        st.font.name = "Times New Roman"
        st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(12)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0,0,0)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(0 if sname=="Heading 1" else 6)
        st.paragraph_format.space_after = Pt(6 if sname=="Heading 1" else 3)

    # COVER: table-based vertical distribution
    cover_sec = doc.sections[0]
    clear_footer(cover_sec)
    cover = doc.add_table(rows=3, cols=1)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.autofit = False
    remove_borders(cover)

    for row, height in zip(cover.rows, [Cm(5.1), Cm(10.3), Cm(7.7)]):
        row.height = height
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    c = cover.cell(0,0)
    c.text = ""
    p = c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run((ident.get("judul") or "").upper()); font_run(r,12,bold=True)

    c = cover.cell(1,0); c.text=""
    if LOGO_PATH.exists():
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Cm(4.25))
    else:
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run("[LOGO UMS]"); font_run(r,12,bold=True)
    for txt,bold,after in [
        ("PROPOSAL SKRIPSI",True,8),
        ("Disusun sebagai salah satu syarat penyusunan Skripsi pada Program Studi",False,0),
        ("Teknik Informatika Fakultas Komunikasi dan Informatika",False,10),
        ("Oleh:",True,3),
        ((ident.get("nama") or "").upper(),True,0),
        (ident.get("nim") or "",True,0)
    ]:
        p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing=1; p.paragraph_format.space_after=Pt(after)
        r=p.add_run(txt); font_run(r,12,bold=bold)

    c=cover.cell(2,0); c.text=""
    for txt in [
        ident.get("programStudi","PROGRAM STUDI TEKNIK INFORMATIKA"),
        ident.get("fakultas","FAKULTAS KOMUNIKASI DAN INFORMATIKA"),
        ident.get("universitas","UNIVERSITAS MUHAMMADIYAH SURAKARTA"),
        str(ident.get("tahun","2026"))
    ]:
        p=c.add_paragraph() if c.paragraphs[0].text else c.paragraphs[0]
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing=1
        r=p.add_run(txt.upper()); font_run(r,12,bold=True)

    # Front matter Roman
    front = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front); clear_footer(front); set_page_number(front,"lowerRoman",1); page_field(front.footer.paragraphs[0])

    add_center(doc,"HALAMAN PERSETUJUAN",True,18)
    add_center(doc,"PROPOSAL SKRIPSI",True,12)
    add_center(doc,(ident.get("judul") or "").upper(),True,16)
    add_center(doc,"Oleh:",False,4)
    add_center(doc,(ident.get("nama") or "").upper(),True,0)
    add_center(doc,ident.get("nim") or "",True,18)
    add_center(doc,"Telah diperiksa dan disetujui untuk diajukan pada tahapan akademik selanjutnya.",False,18)
    pemb = ident.get("pembimbing") or "______________________________"
    nidn = ident.get("nidn") or "__________________"
    add_center(doc,f"Surakarta, __________________ {ident.get('tahun','2026')}\n\nDosen Pembimbing\n\n\n\n{pemb}\nNIDN/NIK. {nidn}")

    doc.add_page_break()
    add_center(doc,"DAFTAR ISI",True,12)
    toc_field(doc.add_paragraph())

    # Main decimal
    main = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main); clear_footer(main); set_page_number(main,"decimal",1); page_field(main.footer.paragraphs[0])

    for chapter_index, chapter in enumerate(data.get("chapters", []), 1):
        chapter_title = (chapter.get("title") or f"Bab {chapter_index}").strip()
        add_heading(doc, f"BAB {chapter_index} {chapter_title.upper()}", 1)
        for section_index, section in enumerate(chapter.get("sections", []), 1):
            section_title = (section.get("title") or f"Subbab {section_index}").strip()
            add_heading(doc, f"{chapter_index}.{section_index} {section_title}", 2)
            for block in section.get("blocks", []):
                add_content_block(doc, block, sources, citation_style, citation_numbers)

    # References
    add_heading(doc,"DAFTAR PUSTAKA",1)
    for ref in refs:
        if not str(ref).strip(): continue
        p=doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Cm(-1)
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_after = Pt(6)
        r=p.add_run(str(ref).strip()); font_run(r)
    formatted_sources = bibliography_sources
    if citation_style == "apa":
        formatted_sources = sorted(
            bibliography_sources,
            key=lambda source: str(source.get("authors") or "").casefold(),
        )
    for source in formatted_sources:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Cm(-1)
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_after = Pt(6)
        number = citation_numbers.get(source.get("id"))
        text = format_reference(source, citation_style, number)
        font_run(p.add_run(text))

    update_fields(doc)
    doc.core_properties.title = ident.get("judul","Proposal Skripsi")
    doc.core_properties.author = ident.get("nama","")
    doc.core_properties.subject = "Proposal Skripsi Teknik Informatika"

    out=BytesIO()
    doc.save(out)
    out.seek(0)
    return out
